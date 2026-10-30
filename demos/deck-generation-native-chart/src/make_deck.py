"""Build a small editable demo deck with native (editable) charts.

Neutral demo data only: table_a with segments A, B, C over periods P1..P6.
Uses the python-pptx default template (no company master, no logos).

Usage: python src/make_deck.py [output_path]
"""
import sys
from pathlib import Path

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

# Demo rows from table_a: (period, segment, value)
PERIODS = ["P1", "P2", "P3", "P4", "P5", "P6"]
SEGMENTS = {
    "A": [120, 135, 128, 150, 162, 170],
    "B": [90, 95, 101, 99, 110, 118],
    "C": [60, 58, 66, 72, 70, 81],
}


def add_title(prs, title, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title
    s.placeholders[1].text = subtitle


def add_bar_chart(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])  # title only
    s.shapes.title.text = "Value by segment"
    data = CategoryChartData()
    data.categories = list(SEGMENTS)
    data.add_series("Total value", [sum(v) for v in SEGMENTS.values()])
    gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6),
                            Inches(8), Inches(5.4), data)
    ch = gf.chart
    ch.has_title = False  # slide title carries the label
    ch.has_legend = False
    plot = ch.plots[0]
    plot.has_data_labels = True
    plot.data_labels.font.size = Pt(14)
    ch.value_axis.has_major_gridlines = True
    ch.category_axis.tick_labels.font.size = Pt(14)


def add_line_chart(prs):
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "Value by period"
    data = CategoryChartData()
    data.categories = PERIODS
    for seg, vals in SEGMENTS.items():
        data.add_series(f"Segment {seg}", vals)
    gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(1), Inches(1.6),
                            Inches(8), Inches(5.4), data)
    ch = gf.chart
    ch.has_legend = True
    ch.legend.position = XL_LEGEND_POSITION.BOTTOM
    ch.legend.include_in_layout = False


def main(out):
    prs = Presentation()  # default blank template
    add_title(prs, "Demo deck", "Generated from table_a (demo data)")
    add_bar_chart(prs)
    add_line_chart(prs)
    cp = prs.core_properties  # default template carries a person name here
    cp.author = cp.last_modified_by = "demo"
    cp.title = "Demo deck"
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    print(f"saved {out}")


if __name__ == "__main__":
    here = Path(__file__).resolve().parent.parent
    main(sys.argv[1] if len(sys.argv) > 1 else here / "output" / "demo_deck.pptx")
